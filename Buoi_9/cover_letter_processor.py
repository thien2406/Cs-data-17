import os
import re

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from docx import Document


class CoverLetterProcessor:

    def __init__(self, folder_path, excel_file):

        self.folder_path = folder_path
        self.excel_file = excel_file

        self.wb = None
        self.ws = None

        # Danh sách các cột trong Excel
        self.headers = [
            "Tên file",
            "Họ và tên",
            "Giới tính",
            "Ngày sinh",
            "Nơi sinh",
            "Nguyên quán",
            "Hộ khẩu thường trú",
            "Chỗ ở hiện nay",
            "Điện thoại"
        ]

        # Các mẫu Regex để tìm thông tin
        self.patterns = {

            "Họ và tên":
                r"Họ và tên\s*:\s*(.*?)\s+Nam/Nữ\s*:",

            "Giới tính":
                r"Nam/Nữ\s*:\s*([^\n]+)",

            "Ngày sinh":
                r"Sinh ngày\s*:\s*(.*?)\s+Nơi sinh\s*:",

            "Nơi sinh":
                r"Nơi sinh\s*:\s*([^\n]+)",

            "Nguyên quán":
                r"Nguyên quán\s*:\s*([^\n]+)",

            "Hộ khẩu thường trú":
                r"Nơi đăng ký hộ khẩu thường trú\s*:\s*([^\n]+)",

            "Chỗ ở hiện nay":
                r"Chỗ ở hiện nay\s*:\s*([^\n]+)",

            "Điện thoại":
                r"Điện thoại(?: liên hệ)?\s*:\s*([^\n]+)"
        }


    def initialize_excel(self):

        if os.path.exists(self.excel_file):

            self.wb = load_workbook(self.excel_file)
            self.ws = self.wb.active

        else:

            self.wb = Workbook()
            self.ws = self.wb.active

            self.ws.title = "Thông tin người dùng"

            self.ws.append(self.headers)


    def read_docx(self, file_path):

        doc = Document(file_path)

        doc_content = [
            paragraph.text
            for paragraph in doc.paragraphs
        ]

        doc_full = "\n".join(doc_content)

        return doc_full


    def extract_info(self, text):

        info = {}

        for key, pattern in self.patterns.items():

            match = re.search(
                pattern,
                text,
                re.IGNORECASE
            )

            if match:

                value = match.group(1).strip()

                # Xóa khoảng trắng thừa
                value = re.sub(
                    r"\s+",
                    " ",
                    value
                )

                info[key] = value

            else:

                # Nếu không tìm thấy
                info[key] = "Không tìm thấy"

        return info


    def auto_adjust_column_width(self):

        for column in self.ws.columns:

            max_length = 0

            # Lấy số thứ tự cột
            column_number = column[0].column

            # Chuyển số thành chữ
            column_letter = get_column_letter(
                column_number
            )

            for cell in column:

                if cell.value is not None:

                    cell_length = len(
                        str(cell.value)
                    )

                    if cell_length > max_length:
                        max_length = cell_length

            # Cộng thêm 2 để nội dung dễ nhìn
            self.ws.column_dimensions[
                column_letter
            ].width = max_length + 2


    def process_documents(self):

        self.initialize_excel()

        doc_files = os.listdir(
            self.folder_path
        )

        # Biến đếm
        success_count = 0
        error_count = 0

        # Duyệt từng file
        for file_name in doc_files:

            # Chỉ xử lý file .docx
            if not file_name.lower().endswith(
                ".docx"
            ):
                continue

            file_path = os.path.join(
                self.folder_path,
                file_name
            )

            try:

                # Đọc Word
                document_text = self.read_docx(
                    file_path
                )

                # Trích xuất thông tin
                data = self.extract_info(
                    document_text
                )

                # Thêm tên file trước
                values = [file_name]

                # Thêm các thông tin còn lại
                for header in self.headers[1:]:

                    values.append(
                        data.get(
                            header,
                            "Không tìm thấy"
                        )
                    )

                # Ghi vào Excel
                self.ws.append(values)

                # Tăng số file thành công
                success_count += 1

                print(
                    f"Đã xử lý thành công: {file_name}"
                )

            except Exception as error:

                # File bị lỗi
                error_count += 1

                print(
                    f"File bị lỗi: {file_name}"
                )

                print(
                    f"Chi tiết lỗi: {error}"
                )

        # Tự động điều chỉnh cột
        self.auto_adjust_column_width()

        # Lưu Excel
        self.wb.save(self.excel_file)

        # In kết quả
        print()
        print("==============================")
        print("KẾT QUẢ XỬ LÝ")
        print("==============================")

        print(
            f"Số file xử lý thành công: {success_count}"
        )

        print(
            f"Số file bị lỗi: {error_count}"
        )

        print(
            f"File Excel: {self.excel_file}"
        )

        print("==============================")
